"""The Word report of a study, with python-docx.

A title page, the answer, the key numbers, every section of ``ws.report``
as a Heading 1 with its Markdown turned into paragraphs (bold, bullets,
numbered lists, inline code), the figures from the PNG artifacts at 6 inches
wide with their captions, the tables as real Word tables (40 rows at most,
with a note when capped), the references, an appendix with ``study.yaml``
in a monospace paragraph, and the footer. python-docx is imported inside
:func:`report_docx_bytes`; when it is missing the function returns None and
leaves a note event on the workspace.
"""

from __future__ import annotations

import io
import re
from typing import Any

from aquascope.studio.deliverables import _common as c
from aquascope.studio.deliverables.tables import frame_of
from aquascope.studio.workspace import Workspace

MAX_TABLE_ROWS = 40
FIGURE_WIDTH_IN = 6.0

_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")
_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_HEADING = re.compile(r"^\s*(#{1,6})\s+(.*)$")


def _runs(paragraph: Any, text: str) -> None:
    """Add ``text`` to a paragraph as runs: ``**bold**``, ``*italic*`` and ```code``` become formatting."""
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _style_or_default(doc: Any, name: str) -> str | None:
    try:
        doc.styles[name]
        return name
    except KeyError:
        return None


def markdown_to_docx(doc: Any, text: str) -> None:
    """A small converter: headings, bullet and numbered lists, paragraphs; inline bold, italic and code."""
    bullet = _style_or_default(doc, "List Bullet")
    numbered = _style_or_default(doc, "List Number")
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            _runs(doc.add_paragraph(), " ".join(buffer))
            buffer.clear()

    for line in text.splitlines():
        if not line.strip():
            flush()
            continue
        m = _HEADING.match(line)
        if m:
            flush()
            doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)) + 1, 4))
            continue
        m = _BULLET.match(line)
        if m:
            flush()
            _runs(doc.add_paragraph(style=bullet) if bullet else doc.add_paragraph(), m.group(1))
            continue
        m = _NUMBERED.match(line)
        if m:
            flush()
            _runs(doc.add_paragraph(style=numbered) if numbered else doc.add_paragraph(), m.group(1))
            continue
        buffer.append(line.strip())
    flush()


def _add_table(doc: Any, columns: list[str], rows: list[list[Any]], caption: str | None) -> None:
    from docx.shared import Pt

    if caption:
        p = doc.add_paragraph()
        p.add_run(caption).italic = True
    table = doc.add_table(rows=1, cols=len(columns))
    style = _style_or_default(doc, "Table Grid")
    if style:
        table.style = style
    for i, col in enumerate(columns):
        cell = table.rows[0].cells[i]
        cell.text = str(col)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
    for r in rows[:MAX_TABLE_ROWS]:
        cells = table.add_row().cells
        for i, v in enumerate(r[: len(columns)]):
            cells[i].text = "" if v is None else str(v)
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(9)
    if len(rows) > MAX_TABLE_ROWS:
        note = doc.add_paragraph()
        note.add_run(f"First {MAX_TABLE_ROWS} of {len(rows)} rows; the full table is in the workbook and the "
                     f"CSV.").italic = True


def _add_figure(doc: Any, png: bytes, caption: str | None) -> None:
    from docx.shared import Inches

    doc.add_picture(io.BytesIO(png), width=Inches(FIGURE_WIDTH_IN))
    if caption:
        p = doc.add_paragraph()
        p.add_run(caption).italic = True


def _monospace(doc: Any, text: str) -> None:
    from docx.shared import Pt

    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(8)


def report_docx_bytes(ws: Workspace) -> bytes | None:
    """``report.docx`` as bytes, or None (with a note event) when python-docx cannot be imported."""
    try:
        from docx import Document
    except ImportError:
        ws.event("author", "note", "python-docx is not installed; report.docx was skipped "
                                   "(pip install python-docx, or aquascope[studio])")
        return None

    doc = Document()
    doc.add_heading(c.title_of(ws), level=0)
    site = c.site_text(ws)
    if site:
        doc.add_paragraph(f"Site: {site}")
    doc.add_paragraph(f"Date: {c.date_of(ws)}")
    doc.add_paragraph("Prepared by AquaScope Studio")
    doc.add_paragraph(c.model_line(ws))
    doc.add_paragraph(f"AquaScope {c.version()}")
    doc.add_page_break()

    answer = c.answer_of(ws)
    if answer:
        doc.add_heading("Answer", level=1)
        _runs(doc.add_paragraph(), answer)
    keys = c.key_numbers(ws)
    if keys:
        _add_table(doc, ["Quantity", "Value", "Unit", "Step"],
                   [[k.get("label"), k.get("value"), k.get("unit") or "", k.get("step") or ""] for k in keys],
                   "Key numbers")

    for block in c.blocks(ws):
        if block["kind"] == "list":
            doc.add_heading(block["title"], level=1)
            if block["key"] == "references":
                markdown_to_docx(doc, "\n".join(f"{i}. {r}" for i, r in enumerate(block["items"], 1)))
            else:
                markdown_to_docx(doc, "\n".join(f"- {i}" for i in block["items"]))
            continue
        doc.add_heading(block["title"] or block["id"], level=1)
        markdown_to_docx(doc, block["text"])
        for fid in block["figures"]:
            fig = c.png_figure(ws, fid)
            if fig is not None and fig.data:
                _add_figure(doc, fig.data, fig.caption)
        for tid in block["tables"]:
            tab = c.table_artifact(ws, tid)
            if tab is None:
                continue
            try:
                columns, values = c.frame_cells(frame_of(tab))
            except Exception:  # noqa: BLE001
                doc.add_paragraph(f"Table {tab.id} could not be rendered.")
                continue
            _add_table(doc, columns, values, tab.caption or tab.id)

    doc.add_heading("Appendix: reproducibility", level=1)
    doc.add_paragraph("The study file below replays with no model: aquascope run study.yaml. The notebook "
                      "study.ipynb in the bundle does the same and redraws the figures.")
    yaml_text = c.study_yaml(ws)
    if yaml_text:
        _monospace(doc, yaml_text)
    if ws.ledger:
        _add_table(doc, ["Role", "Calls", "Prompt tokens", "Completion tokens"],
                   [[role, v.get("calls", 0), v.get("prompt_tokens", 0), v.get("completion_tokens", 0)]
                    for role, v in ws.ledger.items()], "Model calls per role")
    doc.add_paragraph(f"How to cite: {c.citation()}")
    footer = c.footer_of(ws)
    if footer:
        p = doc.add_paragraph()
        p.add_run(footer).italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
